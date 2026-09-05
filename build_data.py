"""Unisce listone Gazzetta (players.json) e file Fantalgoritmo (fantalgoritmo.xlsx) in public/data.json.
Uso: python3 build_data.py    (richiede: pip install openpyxl)
Per aggiornare l'Excel: sostituisci fantalgoritmo.xlsx e rilancia. Genera anche senza_excel.csv (giocatori del listone senza riga Fantalgoritmo)."""
import openpyxl, json, re, unicodedata, collections, csv

L = json.load(open("players.json", encoding="utf-8"))
wb = openpyxl.load_workbook("fantalgoritmo.xlsx", data_only=True)
TEAM = {"atalanta":"ATA","bologna":"BOL","cagliari":"CAG","como":"COM","fiorentina":"FIO","frosinone":"FRO","genoa":"GEN","inter":"INT","juventus":"JUV","lazio":"LAZ","lecce":"LEC","milan":"MIL","monza":"MON","napoli":"NAP","parma":"PAR","roma":"ROM","sassuolo":"SAS","torino":"TOR","udinese":"UDI","venezia":"VEN"}
ROLE = {"Portieri":"POR","Difensori":"DIF","Centrocampisti":"CEN","Attaccanti":"ATT"}
ALLOWED_CROSS = {("CEN","ATT"),("ATT","CEN")}  # ruoli mantra vs Gazzetta: esterni/trequartisti

def norm(s):
    s = unicodedata.normalize("NFD", str(s or "")).encode("ascii","ignore").decode().lower()
    s = re.sub(r"\b[a-z]{1,3}\.\s*", "", s)
    return re.sub(r"[^a-z ]", "", s).strip()
def keys(s):
    n = norm(s); ws = n.split(); out = {n.replace(" ","")}
    if len(ws) > 1: out |= {ws[-1], ws[0], "".join(ws[-2:])}
    return out
def num(v):
    if isinstance(v,(int,float)): return round(float(v),2)
    try: return round(float(str(v).replace(",",".")),2)
    except: return None
def txt(v):
    v = str(v or "").strip(); return v if v and v.lower() not in ("nd","n.d.","0. n.d","0. n.d.","prezzo n.d.","in aggiorn.") else None

# --- Excel rows -> fa objects
ex = []
for ws in wb.worksheets:
    hdr = [str(c.value or "").strip() for c in ws[1]]
    for r in ws.iter_rows(min_row=2, values_only=True):
        d = dict(zip(hdr, r))
        if not d.get("Nome"): continue
        role = ROLE[ws.title]
        pstat = d.get("P. stat. Max") if "P. stat. Max" in d else d.get("P. stat. A")
        fascia = d.get("Fascia") or d.get("Scelta") or d.get("Valore")
        fa = {
            "ia": num(d.get("IA")), "pma": num(d.get("P. Med. Aste")), "pstat": num(pstat), "pgol": num(d.get("P. Gol M.")),
            "fascia": txt(fascia), "note": txt(d.get("Note")), "sos": txt(d.get("SOS")),
            "qt": num(d.get("Qt. Fanta")), "fv": num(d.get("FV")), "trM": num(d.get("Tr. M.")), "trG": num(d.get("Tr. G.")),
            "pg": num(d.get("PG")), "media": num(d.get("Media")), "fmed": num(d.get("F.Med")),
            "gol": num(d.get("Gol")), "ass": num(d.get("Ass.")), "amm": num(d.get("Amm.")), "golSub": num(d.get("Gol subiti")),
            "mantra": txt(d.get("Ruolo")), "accoppiata": txt(d.get("Accoppiata squadra migliore") or d.get("Accoppiata Sq. F.-D.")),
            "rank": d.get("I."),
        }
        fa = {k:v for k,v in fa.items() if v is not None}
        ex.append({"nome": str(d["Nome"]).strip(), "team": TEAM.get(str(d.get("Squadra") or "").strip().lower()), "squadra": txt(d.get("Squadra")), "role": role, "fa": fa})

# --- matching
byname = collections.defaultdict(list)
for i,p in enumerate(L):
    for k in keys(p["name"]): byname[k].append(i)
used, res = set(), {}
def try_match(d, strict):
    cands = set()
    for k in keys(d["nome"]): cands |= set(byname.get(k, []))
    # ruolo diverso accettato solo se coppia mantra/Gazzetta nota, oppure stessa squadra (es. terzino classificato C)
    cands = [i for i in cands if i not in used and (L[i]["role"] == d["role"] or (d["role"], L[i]["role"]) in ALLOWED_CROSS or (d["team"] and L[i]["team"] == d["team"]))]
    lvls = [[i for i in cands if L[i]["team"] == d["team"] and L[i]["role"] == d["role"]],
            [i for i in cands if L[i]["role"] == d["role"]],
            [i for i in cands if L[i]["team"] == d["team"]],
            cands]  # ultimo livello: solo nome, accettato solo se unico (e ruolo compatibile)
    for n,c in enumerate(lvls):
        if strict and n > 0: break
        if len(c) == 1: return c[0]
        if len(c) > 1: return None
    return None
for strict in (True, False):
    for j,d in enumerate(ex):
        if j in res: continue
        i = try_match(d, strict)
        if i is not None: res[j] = i; used.add(i)

# --- output dataset
out = []
for i,p in enumerate(L):
    out.append({"id": f"{p['name']}|{p['team']}|{p['role']}", "name": p["name"], "team": p["team"], "role": p["role"], "cost": p["cost"], "src": "listone", "fa": None})
for j,d in enumerate(ex):
    if j in res: out[res[j]]["fa"] = d["fa"]; out[res[j]]["faName"] = d["nome"]
    else:
        out.append({"id": f"{d['nome'].upper()}|{d['team'] or '---'}|{d['role']}", "name": d["nome"].upper(), "team": d["team"] or "---", "role": d["role"], "cost": None, "src": "excel", "fa": d["fa"], "faName": d["nome"]})
json.dump({"base": 1000, "players": out}, open("public/data.json","w",encoding="utf-8"), ensure_ascii=False, separators=(",",":"))

# --- report
nol = [p for i,p in enumerate(L) if i not in used]
with open("senza_excel.csv","w",newline="",encoding="utf-8") as f:
    w = csv.writer(f, delimiter=";"); w.writerow(["Nome","Squadra","Ruolo","Costo listone"])
    for p in sorted(nol, key=lambda p:(p["role"], -p["cost"])): w.writerow([p["name"],p["team"],p["role"],p["cost"]])
with open("senza_listone.csv","w",newline="",encoding="utf-8") as f:
    w = csv.writer(f, delimiter=";"); w.writerow(["Nome Excel","Squadra","Ruolo","IA","P. Med. Aste"])
    for j,d in enumerate(ex):
        if j not in res: w.writerow([d["nome"], d["squadra"], d["role"], d["fa"].get("ia"), d["fa"].get("pma")])
print(f"listone {len(L)} · excel {len(ex)} · abbinati {len(res)} · listone senza excel {len(nol)} · excel senza listone {len(ex)-len(res)} · dataset {len(out)}")
